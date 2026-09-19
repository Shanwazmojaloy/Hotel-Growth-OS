"""Tests for the prospect screener.

The property that matters: the screener must refuse a confident answer when the
verdict depends on an assumption the prospect has not confirmed. In testing, a
single input moved a real prospect from NOT_VIABLE to VIABLE — an overconfident
call list is worse than no call list.
"""
from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from prospect.review import slugify
from prospect.screener import (
    Prospect, ProspectError, format_table, load_prospects, screen, screen_all, to_csv,
)

DATA = Path(__file__).resolve().parents[1] / "data"

FEE, MEDIA = 19_000.0, 50_000.0


def prospect(**kw) -> Prospect:
    base = dict(name="Test Hotel", city="Dhaka", segment="business", rooms=60,
                adr_bdt=9_000, occupancy=0.60, ota_share=0.65,
                effective_ota_commission=0.17, direct_share=0.20,
                adr_basis="listed", rooms_status="published", assumed_fields=[], source_id="TEST")
    base.update(kw)
    return Prospect(**base)


class TestDerivations(unittest.TestCase):
    def test_room_nights_and_commission_are_computed(self):
        p = screen(prospect(), fee_bdt=FEE, media_bdt=MEDIA)
        # 60 rooms × 30.44 days × 0.60 occupancy
        self.assertAlmostEqual(p.total_room_nights, 60 * 30.44 * 0.60, places=1)
        self.assertAlmostEqual(p.ota_room_nights, p.total_room_nights * 0.65, places=1)
        self.assertAlmostEqual(p.commission_bdt,
                               p.ota_room_nights * 9_000 * 0.17, places=0)

    def test_saving_per_room_night_uses_the_direct_cost_gap(self):
        p = screen(prospect(), fee_bdt=FEE, media_bdt=MEDIA)
        self.assertAlmostEqual(p.saving_per_room_night, 9_000 * (0.17 - 0.045), places=2)

    def test_capturable_uses_the_conservative_shift_share(self):
        """The pitch number must be the bottom of the plausible range."""
        from prospect.screener import CONSERVATIVE_SHIFTABLE_SHARE
        p = screen(prospect(), fee_bdt=FEE, media_bdt=MEDIA)
        expected = p.ota_room_nights * CONSERVATIVE_SHIFTABLE_SHARE * p.saving_per_room_night
        self.assertAlmostEqual(p.capturable_bdt, expected, places=0)
        self.assertLess(CONSERVATIVE_SHIFTABLE_SHARE, 0.30)

    def test_direct_share_is_derived_and_flagged_when_missing(self):
        p = screen(prospect(direct_share=None), fee_bdt=FEE, media_bdt=MEDIA)
        self.assertTrue(p.derived_direct_share)
        self.assertTrue(any("direct volume was estimated" in r for r in p.reasons))

    def test_stated_direct_share_is_not_flagged(self):
        p = screen(prospect(direct_share=0.25), fee_bdt=FEE, media_bdt=MEDIA)
        self.assertFalse(p.derived_direct_share)
        self.assertFalse(any("estimated" in r for r in p.reasons))


class TestRecommendations(unittest.TestCase):
    def test_a_large_healthy_hotel_is_called_now(self):
        p = screen(prospect(rooms=120, adr_bdt=18_000, occupancy=0.70, ota_share=0.40,
                            effective_ota_commission=0.15, direct_share=0.28),
                   fee_bdt=FEE, media_bdt=MEDIA)
        self.assertEqual(p.recommendation, "CALL NOW")
        self.assertLess(p.lift_pessimistic, 0.35)

    def test_a_thin_small_hotel_is_not_a_fit(self):
        p = screen(prospect(rooms=30, adr_bdt=6_000, occupancy=0.51, ota_share=0.77,
                            effective_ota_commission=0.15),
                   fee_bdt=FEE, media_bdt=MEDIA)
        self.assertEqual(p.recommendation, "NOT A FIT")
        self.assertGreater(p.lift_optimistic, 0.60)

    def test_confirm_state_when_the_verdict_depends_on_the_assumption(self):
        """The core honesty property: if a favourable and a pessimistic reading
        give different verdicts, the screener says so instead of guessing."""
        p = screen(prospect(rooms=72, adr_bdt=11_000, occupancy=0.51, ota_share=0.73,
                            effective_ota_commission=0.18, direct_share=None),
                   fee_bdt=FEE, media_bdt=MEDIA)
        self.assertEqual(p.recommendation, "CONFIRM DIRECT VOLUME FIRST")
        self.assertLess(p.lift_optimistic, 0.35)
        self.assertGreater(p.lift_pessimistic, 0.35)
        self.assertTrue(any("how many bookings arrive by phone" in r for r in p.reasons))

    def test_a_tiny_capturable_pool_is_not_worth_a_call(self):
        p = screen(prospect(rooms=10, adr_bdt=5_000, occupancy=0.5, ota_share=0.6,
                            effective_ota_commission=0.15, direct_share=0.30),
                   fee_bdt=FEE, media_bdt=MEDIA)
        self.assertIn(p.recommendation, {"CALL LATER", "NOT A FIT"})

    def test_illiquid_economics_are_never_called_now(self):
        """Zero commission gap → nothing to save → never a call-now."""
        p = screen(prospect(effective_ota_commission=0.045), fee_bdt=FEE, media_bdt=MEDIA)
        self.assertNotEqual(p.recommendation, "CALL NOW")


class TestRanking(unittest.TestCase):
    def test_ranking_puts_call_now_first_by_opportunity(self):
        prospects = [
            prospect(name="Small", rooms=30, adr_bdt=6_000, occupancy=0.51, ota_share=0.77),
            prospect(name="Big", rooms=120, adr_bdt=18_000, occupancy=0.70, ota_share=0.40,
                     effective_ota_commission=0.15, direct_share=0.28),
            prospect(name="Medium", rooms=85, adr_bdt=9_500, occupancy=0.66, ota_share=0.52),
        ]
        ranked = screen_all(prospects, fee_bdt=FEE, media_bdt=MEDIA)
        order = [p.recommendation for p in ranked]
        self.assertEqual(order, sorted(order, key=lambda r: ["CALL NOW", "CONFIRM DIRECT VOLUME FIRST",
                                                             "CALL LATER", "NOT A FIT"].index(r)))
        call_now = [p for p in ranked if p.recommendation == "CALL NOW"]
        self.assertEqual(call_now[0].name, "Big")

    def test_table_renders_without_error(self):
        ranked = screen_all([prospect(), prospect(name="B", rooms=30, adr_bdt=6000)],
                            fee_bdt=FEE, media_bdt=MEDIA)
        text = format_table(ranked)
        self.assertIn("RECOMMENDATION", text)
        # header + separator + one row per prospect + separator
        self.assertEqual(len(text.splitlines()), len(ranked) + 3)


class TestInputValidation(unittest.TestCase):
    def test_missing_column_is_a_hard_error(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "x.csv"
            p.write_text("name,city,rooms\nA,Dhaka,50\n")
            with self.assertRaises(ProspectError) as ctx:
                load_prospects(p)
            self.assertIn("missing columns", str(ctx.exception))

    def test_non_numeric_is_a_hard_error(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "x.csv"
            p.write_text("name,city,segment,rooms,adr_bdt,occupancy,ota_share,"
                         "effective_ota_commission\nA,Dhaka,business,many,9000,0.6,0.5,0.16\n")
            with self.assertRaises(ProspectError):
                load_prospects(p)

    def test_occupancy_must_be_a_fraction_not_a_percentage(self):
        p = prospect(occupancy=60.0)
        with self.assertRaises(ProspectError) as ctx:
            screen(p, fee_bdt=FEE, media_bdt=MEDIA)
        self.assertIn("fraction", str(ctx.exception))

    def test_round_trip_through_csv(self):
        ranked = screen_all([prospect()], fee_bdt=FEE, media_bdt=MEDIA)
        with tempfile.TemporaryDirectory() as td:
            out = to_csv(ranked, Path(td) / "s.csv")
            rows = load_prospects(out)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].name, "Test Hotel")


class TestSampleFile(unittest.TestCase):
    def test_sample_prospects_load_and_rank(self):
        path = DATA / "prospects_sample.csv"
        if not path.exists():
            self.skipTest("sample prospect file missing")
        ranked = screen_all(load_prospects(path), fee_bdt=FEE, media_bdt=MEDIA)
        self.assertGreaterEqual(len(ranked), 5)
        self.assertTrue(any(p.recommendation == "CALL NOW" for p in ranked))
        self.assertTrue(any(p.recommendation == "NOT A FIT" for p in ranked))
        # a call list that says yes to everything is not a call list
        call_now = [p for p in ranked if p.recommendation == "CALL NOW"]
        self.assertLess(len(call_now), len(ranked))

    def test_slugify_produces_usable_filenames(self):
        self.assertEqual(slugify("Gulshan business hotel — 120 rooms (archetype B)"),
                         "gulshan_business_hotel_120_rooms")
        self.assertEqual(slugify("!!!"), "prospect")


if __name__ == "__main__":
    unittest.main(verbosity=2)


class TestHonestyGate(unittest.TestCase):
    """A gate that cannot fail is decoration. These mutations prove it can."""

    @classmethod
    def setUpClass(cls):
        from prospect.screener import screen_all, load_prospects
        cls.scores = [p.as_dict() for p in
                      screen_all(load_prospects(DATA / "prospects_sample.csv"),
                                 fee_bdt=FEE, media_bdt=MEDIA)]

    def _expect_failure(self, mutate, needle: str):
        from prospect.honesty_gate import GateFailure, check
        scores = json.loads(json.dumps(self.scores))  # deep copy
        mutate(scores)
        with self.assertRaises(GateFailure) as ctx:
            check(scores)
        self.assertIn(needle, str(ctx.exception))

    def test_unmutated_output_passes(self):
        from prospect.honesty_gate import check
        self.assertTrue(check(self.scores))

    def test_gate_rejects_a_confident_yes_on_an_estimated_volume(self):
        self._expect_failure(
            lambda s: next(p for p in s if p["recommendation"] == "CONFIRM DIRECT VOLUME FIRST"
                           ).update(recommendation="CALL NOW"), "rule gives")

    def test_gate_rejects_a_hopeless_prospect_presented_as_workable(self):
        def mutate(s):
            p = next(p for p in s if p["recommendation"] == "NOT A FIT")
            p.update(recommendation="CONFIRM DIRECT VOLUME FIRST")
        self._expect_failure(mutate, "rule gives")

    def test_gate_rejects_an_inflated_quote(self):
        def mutate(s):
            p = next(p for p in s if p["recommendation"] == "CALL NOW")
            p["capturable_bdt"] = p["commission_bdt"] * 0.9
        self._expect_failure(mutate, "as capturable")

    def test_gate_rejects_an_inverted_band(self):
        def mutate(s):
            p = s[0]
            p["lift_optimistic"], p["lift_pessimistic"] = p["lift_pessimistic"], p["lift_optimistic"]
        self._expect_failure(mutate, "inverted band")

    def test_gate_rejects_an_empty_list(self):
        from prospect.honesty_gate import GateFailure, check
        with self.assertRaises(GateFailure):
            check([])


class TestAssumptionDiscipline(unittest.TestCase):
    """A confident yes must never rest on a figure we assumed."""

    def test_assumed_inputs_cap_the_verdict_at_confirm(self):
        p = prospect(rooms=120, adr_bdt=18_000, occupancy=0.70, ota_share=0.40,
                     effective_ota_commission=0.15, direct_share=0.28,
                     assumed_fields=["occupancy"])
        scored = screen(p, fee_bdt=FEE, media_bdt=MEDIA)
        self.assertEqual(scored.recommendation, "CONFIRM DIRECT VOLUME FIRST")
        self.assertTrue(any("assumptions" in r for r in scored.reasons))

    def test_same_economics_without_assumptions_is_a_call(self):
        p = prospect(rooms=120, adr_bdt=18_000, occupancy=0.70, ota_share=0.40,
                     effective_ota_commission=0.15, direct_share=0.28)
        self.assertEqual(screen(p, fee_bdt=FEE, media_bdt=MEDIA).recommendation, "CALL NOW")

    def test_provenance_columns_survive_a_round_trip(self):
        p = prospect(assumed_fields=["occupancy", "ota_share"], adr_basis="class_avg",
                     rooms_status="disputed", source_id="S9")
        scored = screen(p, fee_bdt=FEE, media_bdt=MEDIA)
        d = scored.as_dict()
        self.assertEqual(d["assumed_fields"], "occupancy;ota_share")
        self.assertEqual(d["adr_basis"], "class_avg")
        self.assertEqual(d["rooms_status"], "disputed")

    def test_dhaka_list_loads_with_provenance(self):
        path = DATA / "dhaka_2026-09.csv"
        if not path.exists():
            self.skipTest("Dhaka list missing")
        rows = [screen(p, fee_bdt=FEE, media_bdt=MEDIA) for p in load_prospects(path)]
        self.assertGreaterEqual(len(rows), 12)
        # every row must declare what is assumed, and none may claim a confident call
        for r in rows:
            self.assertTrue(r.assumed_fields, f"{r.name} declares no assumptions — public data cannot be complete")
            self.assertNotEqual(r.recommendation, "CALL NOW", f"{r.name} claimed a confident call")
            self.assertTrue(r.source_id, f"{r.name} has no source")
        # and the screen must still discriminate on a real list
        verdicts = {r.recommendation for r in rows}
        self.assertGreaterEqual(len(verdicts), 2)
        self.assertIn("NOT A FIT", verdicts)


class TestCallSheet(unittest.TestCase):
    def _sheet(self) -> str:
        from prospect.callsheet import build_callsheet
        rows = screen_all(load_prospects(DATA / "dhaka_2026-09.csv"), fee_bdt=FEE, media_bdt=MEDIA)
        return build_callsheet(rows, fee_bdt=FEE, media_bdt=MEDIA, city="Dhaka",
                               source_file="dhaka_2026-09.csv")

    def test_sheet_contains_the_working_sections(self):
        text = self._sheet()
        for heading in ("## Work these, in order", "## Parked (CALL LATER)",
                        "## Screened out (NOT A FIT)", "## What the structure costs you",
                        "## Before you quote anything from this sheet"):
            self.assertIn(heading, text)

    def test_sheet_never_quotes_an_unconfirmed_number_as_fact(self):
        text = self._sheet()
        # the lead-with line must always be paired with what to confirm
        self.assertEqual(text.count("**Lead with:**"), text.count("**Confirm on the call before quoting anything:**"))

    def test_sensitivity_table_uses_the_real_media_budget(self):
        text = self._sheet()
        self.assertIn("৳0 (fee only)", text)
        self.assertIn("৳50,000", text)


class TestGateMirrorsTheRule(unittest.TestCase):
    """expected_verdict() is a restatement of screener.decide(). It must be exact.

    When it drifted, the gate rejected CORRECT output — it demanded CALL NOW for a row
    whose assumptions also forbid CALL NOW — and the cheapest way to satisfy a gate like
    that is to weaken the screener. A checker that fails correct work is worse than no
    checker, so the mirror is pinned here on every branch.
    """

    def _dict(self, **kw) -> dict:
        return screen(prospect(**kw), fee_bdt=FEE, media_bdt=MEDIA).as_dict()

    def test_mirror_holds_on_every_branch(self):
        from prospect.honesty_gate import expected_verdict
        cases = {
            "clean call": dict(rooms=200, adr_bdt=20_000, occupancy=0.75, ota_share=0.60,
                               effective_ota_commission=0.20, direct_share=0.25),
            "stated volume, assumed inputs": dict(rooms=200, adr_bdt=20_000, occupancy=0.75,
                                                  ota_share=0.60, effective_ota_commission=0.20,
                                                  direct_share=0.25,
                                                  assumed_fields=["occupancy", "ota_share"]),
            "volume never stated": dict(rooms=200, adr_bdt=20_000, occupancy=0.75, ota_share=0.60,
                                        effective_ota_commission=0.20, direct_share=None),
            "capturable below the floor": dict(rooms=30, adr_bdt=12_000, occupancy=0.50,
                                               ota_share=0.15, effective_ota_commission=0.18,
                                               direct_share=0.55),
            "hopeless": dict(rooms=8, adr_bdt=2_500, occupancy=0.40, ota_share=0.30,
                             effective_ota_commission=0.15, direct_share=0.30),
            "seasonal downgrade": dict(rooms=50, adr_bdt=6_000, occupancy=0.50, ota_share=0.35,
                                       effective_ota_commission=0.18, direct_share=0.56,
                                       season_ratio=6.0),
            "seasonal but survivable": dict(season_ratio=2.0),
        }
        for label, kw in cases.items():
            with self.subTest(branch=label):
                d = self._dict(**kw)
                self.assertEqual(expected_verdict(d), d["recommendation"])

    def test_the_two_regression_shapes_land_where_the_discipline_says(self):
        assumed = self._dict(rooms=200, adr_bdt=20_000, occupancy=0.75, ota_share=0.60,
                             effective_ota_commission=0.20, direct_share=0.25,
                             assumed_fields=["occupancy", "ota_share"])
        self.assertEqual(assumed["recommendation"], "CONFIRM DIRECT VOLUME FIRST")
        seasonal = self._dict(rooms=50, adr_bdt=6_000, occupancy=0.50, ota_share=0.35,
                              effective_ota_commission=0.18, direct_share=0.56, season_ratio=6.0)
        self.assertEqual(seasonal["recommendation"], "CALL LATER")

    def test_gate_passes_a_list_mixing_assumed_and_seasonal_rows(self):
        """The false-positive direction. Before the fix this raised GateFailure."""
        from prospect.honesty_gate import check
        scores = [
            self._dict(rooms=200, adr_bdt=20_000, occupancy=0.75, ota_share=0.60,
                       effective_ota_commission=0.20, direct_share=0.25),
            self._dict(rooms=200, adr_bdt=20_000, occupancy=0.75, ota_share=0.60,
                       effective_ota_commission=0.20, direct_share=0.25,
                       assumed_fields=["occupancy", "ota_share"]),
            self._dict(rooms=50, adr_bdt=6_000, occupancy=0.50, ota_share=0.35,
                       effective_ota_commission=0.18, direct_share=0.56, season_ratio=6.0),
            self._dict(rooms=8, adr_bdt=2_500, occupancy=0.40, ota_share=0.30,
                       effective_ota_commission=0.15, direct_share=0.30),
        ]
        self.assertTrue(check(scores))

    def test_gate_still_rejects_a_confident_yes_on_assumed_inputs(self):
        """The dangerous direction must survive the fix."""
        from prospect.honesty_gate import GateFailure, check
        scores = [
            self._dict(rooms=200, adr_bdt=20_000, occupancy=0.75, ota_share=0.60,
                       effective_ota_commission=0.20, direct_share=0.25,
                       assumed_fields=["occupancy", "ota_share"]),
            self._dict(rooms=8, adr_bdt=2_500, occupancy=0.40, ota_share=0.30,
                       effective_ota_commission=0.15, direct_share=0.30),
        ]
        scores[0]["recommendation"] = "CALL NOW"
        with self.assertRaises(GateFailure):
            check(scores)


class TestSeasonalReachability(unittest.TestCase):
    """The seasonal guard must be reachable from a real prospect file, not only from a
    hand-built dataclass — otherwise it is decoration."""

    def test_season_ratio_reaches_the_screener_through_the_csv_path(self):
        csv_text = (
            "name,city,segment,rooms,adr_bdt,occupancy,ota_share,effective_ota_commission,"
            "direct_share,season_ratio,adr_basis,rooms_status,assumed_fields,source_id,notes\n"
            '"Seasonal Hotel","Cox\'s Bazar","midscale",50,6000,0.50,0.35,0.18,0.56,6.0,'
            'listed,published,"","T1","peak/trough ratio stated by the owner"\n'
        )
        with tempfile.TemporaryDirectory() as td:
            f = Path(td) / "one.csv"
            f.write_text(csv_text, encoding="utf-8")
            scored = screen(load_prospects(f)[0], fee_bdt=FEE, media_bdt=MEDIA)
        self.assertEqual(scored.season_ratio, 6.0)
        self.assertEqual(scored.recommendation, "CALL LATER")
        self.assertTrue(any("season-weighted" in r for r in scored.reasons))

    def test_a_missing_season_column_reads_as_flat_not_as_seasonal(self):
        csv_text = (
            "name,city,segment,rooms,adr_bdt,occupancy,ota_share,effective_ota_commission,"
            "direct_share,adr_basis,rooms_status,assumed_fields,source_id,notes\n"
            '"Flat Hotel","Dhaka","business",50,6000,0.50,0.35,0.18,0.56,'
            'listed,published,"","T1","no seasonality stated"\n'
        )
        with tempfile.TemporaryDirectory() as td:
            f = Path(td) / "one.csv"
            f.write_text(csv_text, encoding="utf-8")
            scored = screen(load_prospects(f)[0], fee_bdt=FEE, media_bdt=MEDIA)
        self.assertEqual(scored.season_ratio, 1.0)
        self.assertNotEqual(scored.recommendation, "CALL LATER")  # flat: not downgraded


class TestOnePagerSeasonality(unittest.TestCase):
    """The document that reaches the prospect must not quote a flat month to a
    seasonal property without saying so."""

    def _text(self, **kw) -> str:
        from docx import Document
        from prospect.review import build_review
        p = screen(prospect(**kw), fee_bdt=FEE, media_bdt=MEDIA)
        with tempfile.TemporaryDirectory() as td:
            path = build_review(p, Path(td) / "review.docx", fee_bdt=FEE, media_bdt=MEDIA)
            doc = Document(str(path))
        parts = [par.text for par in doc.paragraphs]
        for table in doc.tables:                      # caveats live in table cells too
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)

    def test_a_seasonal_property_is_told_so(self):
        text = self._text(rooms=50, adr_bdt=6_000, occupancy=0.50, ota_share=0.35,
                          effective_ota_commission=0.18, direct_share=0.56, season_ratio=6.0)
        self.assertIn("seasonality", text.lower())
        self.assertIn("season-weighted", text)
        self.assertIn("peak-window reading", text)

    def test_a_flat_property_gets_no_seasonality_section(self):
        text = self._text()
        self.assertNotIn("season-weighted", text)
        self.assertNotIn("peak-window reading", text)

    def test_a_seasonal_property_is_never_told_the_flat_story_is_achievable(self):
        """Two claims the owner can hold side by side must not contradict each other."""
        text = self._text(rooms=50, adr_bdt=6_000, occupancy=0.50, ota_share=0.35,
                          effective_ota_commission=0.18, direct_share=0.56, season_ratio=6.0)
        self.assertIn("season-weighted", text)
        self.assertNotIn("which is why I am suggesting we talk", text)
