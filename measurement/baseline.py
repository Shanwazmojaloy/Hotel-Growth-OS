"""Auto-build the measurement plan §2 baseline table from the exports.

This is where three of the fourteen onboarding hours go, by hand, every time:
occupancy, ADR, RevPAR, channel mix, effective commission, seasonality. All of
it is computable from files the hotel already has.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

from .config import HotelConfig


@dataclass
class Baseline:
    window: str
    months_of_history: int
    total_room_nights: float
    available_room_nights: float
    occupancy: float
    adr_bdt: float
    revpar_bdt: float
    total_room_revenue_bdt: float
    ota_room_nights: float
    ota_share: float
    ota_gross_bdt: float
    ota_commission_bdt: float
    effective_ota_commission: float
    direct_room_nights: float
    direct_share: float
    direct_revenue_bdt: float
    direct_adr_bdt: float
    direct_room_nights_per_month: float
    direct_room_nights_per_room_per_month: float
    seasonality: dict[str, float]
    peak_months: list[str]
    dead_months: list[str]
    gaps: list[str]

    def as_dict(self) -> dict:
        return {k: (round(v, 3) if isinstance(v, float) else v) for k, v in asdict(self).items()}


def _days_in(month: str) -> int:
    y, m = (int(x) for x in month.split("-"))
    ny, nm = (y + (m == 12), 1 if m == 12 else m + 1)
    return (date(ny, nm, 1) - date(y, m, 1)).days


def build_baseline(cfg: HotelConfig, loaded: dict, *, exclude_pilot_months: bool = True) -> Baseline:
    pms = loaded["pms"]
    ota = loaded["ota"]
    bookings = loaded["bookings"]

    months = sorted({p["month"] for p in pms})
    if exclude_pilot_months:
        months = [m for m in months if m < cfg.pilot_start]
    if not months:
        months = sorted({p["month"] for p in pms})

    pms_rows = [p for p in pms if p["month"] in months]
    total_rn = sum(p["total_room_nights"] for p in pms_rows)
    total_rev = sum(p["total_revenue_bdt"] for p in pms_rows)
    available = cfg.rooms * sum(_days_in(m) for m in months)

    ota_rows = [o for o in ota if o["month"] in months]
    ota_rn = sum(o["room_nights"] for o in ota_rows)
    ota_gross = sum(o["gross_bdt"] for o in ota_rows)
    ota_comm = sum(o["commission_bdt"] for o in ota_rows)

    direct_rows = [b for b in bookings
                   if f"{b['date']:%Y-%m}" in months and b["channel"] in {"direct", "phone"}]
    direct_rn = sum(b["room_nights"] for b in direct_rows)
    direct_rev = sum(b["revenue_bdt"] for b in direct_rows)

    monthly_rn = {}
    for p in pms_rows:
        monthly_rn[p["month"]] = p["total_room_nights"]
    peak = [m for m, v in monthly_rn.items() if v > 1.15 * total_rn / max(1, len(months))]
    dead = [m for m, v in monthly_rn.items() if v < 0.75 * total_rn / max(1, len(months))]

    # expected months, for gap detection
    expected = []
    if months:
        y, m = (int(x) for x in months[0].split("-"))
        while f"{y:04d}-{m:02d}" <= months[-1]:
            expected.append(f"{y:04d}-{m:02d}")
            m += 1
            if m > 12:
                y, m = y + 1, 1
    gaps = [m for m in expected if m not in months]

    n = max(1, len(months))
    return Baseline(
        window=f"{months[0]} → {months[-1]}",
        months_of_history=len(months),
        total_room_nights=total_rn,
        available_room_nights=available,
        occupancy=total_rn / available if available else 0.0,
        adr_bdt=total_rev / total_rn if total_rn else 0.0,
        revpar_bdt=total_rev / available if available else 0.0,
        total_room_revenue_bdt=total_rev,
        ota_room_nights=ota_rn,
        ota_share=ota_rn / total_rn if total_rn else 0.0,
        ota_gross_bdt=ota_gross,
        ota_commission_bdt=ota_comm,
        effective_ota_commission=ota_comm / ota_gross if ota_gross else 0.0,
        direct_room_nights=direct_rn,
        direct_share=direct_rn / total_rn if total_rn else 0.0,
        direct_revenue_bdt=direct_rev,
        direct_adr_bdt=direct_rev / direct_rn if direct_rn else 0.0,
        direct_room_nights_per_month=direct_rn / n,
        direct_room_nights_per_room_per_month=(direct_rn / n) / cfg.rooms if cfg.rooms else 0.0,
        seasonality={m: round(v / (total_rn / n), 2) for m, v in sorted(monthly_rn.items())},
        peak_months=peak,
        dead_months=dead,
        gaps=gaps,
    )


def build_docx(cfg: HotelConfig, b: Baseline, out_path: Path, pms_rows: list[dict] | None = None) -> Path:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.7)
        s.left_margin = s.right_margin = Cm(1.9)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(5)

    def H(text, level=1):
        h = doc.add_heading(text, level=level)
        for r in h.runs:
            r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    def T(rows, widths=None):
        t = doc.add_table(rows=0, cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            cells = t.add_row().cells
            for j, val in enumerate(row):
                cells[j].text = ""
                run = cells[j].paragraphs[0].add_run(str(val))
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True
        if widths:
            for j, w in enumerate(widths):
                for row in t.rows:
                    row.cells[j].width = Cm(w)
        doc.add_paragraph()

    h = doc.add_heading(f"Baseline — {cfg.property_name}", level=0)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    p = doc.add_paragraph()
    r = p.add_run(
        f"{cfg.city} · {cfg.rooms} rooms · window {b.window} ({b.months_of_history} months) · "
        f"generated from the property's own exports"
    )
    r.italic = True
    r.font.size = Pt(9)

    doc.add_paragraph(
        "Every figure below is computed from the property's exports. §2 of the signed measurement plan is "
        "completed from this table; anything the property disputes is corrected here, before the plan is signed."
    )

    H("Commercial baseline")
    T([
        ["Measure", "Value", "Note"],
        ["Occupancy", f"{b.occupancy:.1%}", f"{b.total_room_nights:,.0f} of {b.available_room_nights:,.0f} available room-nights"],
        ["ADR", f"৳{b.adr_bdt:,.0f}", "Total room revenue ÷ room-nights, PMS basis"],
        ["RevPAR", f"৳{b.revpar_bdt:,.0f}", "The number that actually moves when channel mix improves"],
        ["Total room revenue", f"৳{b.total_room_revenue_bdt:,.0f}", ""],
        ["OTA share of room-nights", f"{b.ota_share:.1%}", f"{b.ota_room_nights:,.0f} room-nights"],
        ["Effective OTA commission", f"{b.effective_ota_commission:.1%}",
         "Total commission ÷ total OTA gross, per statements — NOT the headline rate"],
        ["Commission paid", f"৳{b.ota_commission_bdt:,.0f}", "The money this engagement is trying to reduce"],
        ["Direct share of room-nights", f"{b.direct_share:.1%}", f"{b.direct_room_nights:,.0f} room-nights"],
        ["Direct ADR", f"৳{b.direct_adr_bdt:,.0f}", "Compare with OTA ADR — a large gap usually means discounting"],
        ["Direct volume", f"{b.direct_room_nights_per_month:,.1f}/month",
         f"{b.direct_room_nights_per_room_per_month:.2f} per room per month"],
    ], widths=[5.2, 3.6, 7.2])

    H("Seasonality")
    T([["Month", "Room-nights", "Index vs average"]] +
      [[m, f"{v:,.0f}", f"{b.seasonality[m]:.2f}×"] for m, v in
       sorted({p["month"]: p["total_room_nights"] for p in (pms_rows or [])
               if p["month"] in b.seasonality}.items())],
      widths=[3.5, 4.0, 4.0])
    doc.add_paragraph(
        f"Peak: {', '.join(b.peak_months) or 'none identified'}. "
        f"Dead: {', '.join(b.dead_months) or 'none identified'}. "
        "This is why a flat monthly fee fails in a dead season — agree the off-peak arrangement before signing."
    )

    if b.gaps:
        H("Gaps in the record")
        doc.add_paragraph(
            f"Missing months in the PMS history: {', '.join(b.gaps)}. Confirm with the property whether those "
            "months had no activity or the export is incomplete."
        )

    doc.add_paragraph()
    doc.add_paragraph("Reproduce: python3 -m measurement.cli baseline --hotel <dir> --out reports/").italic = True

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
